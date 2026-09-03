"""
Helpers communs aux tests de non-régression PDF/Word (Batch 11 T3).

Construisent un Devis minimal et extraient le texte réellement rendu par
pdf_service.py / word_service.py, pour vérifier qu'aucune ligne n'est
invisible (bug récurrent du set_text_color persistant de fpdf2, cf. CLAUDE.md)
et que les règles métier (au réel, art. 293 B, signature masquée...)
produisent bien le texte attendu dans le document final.
"""
import io

from app.models.quote import (
    ArtisanInfo,
    ChantierInfo,
    ClientInfo,
    Devis,
    LigneDevis,
    TotauxDevis,
)
from app.services.pdf_service import generate_quote_pdf
from app.services.word_service import generate_quote_docx
from docx import Document
from pypdf import PdfReader


def make_ligne(
    poste: str = "Poste",
    description: str = "Description",
    lot: str | None = None,
    quantite: float | None = 1.0,
    unite: str = "u",
    prix_unitaire_ht: float = 100.0,
    tva_taux: float = 20.0,
) -> LigneDevis:
    return LigneDevis(
        lot=lot,
        poste=poste,
        description=description,
        quantite=quantite,
        unite=unite,
        prix_unitaire_ht=prix_unitaire_ht,
        tva_taux=tva_taux,
    )


def make_devis(lignes: list[LigneDevis] | None = None, **overrides) -> Devis:
    """Devis minimal valide. Passer `lignes=[...]` ou tout champ Devis en override."""
    if lignes is None:
        lignes = [make_ligne()]
    total_ht = round(sum((l.quantite if l.quantite is not None else 1.0) * l.prix_unitaire_ht for l in lignes), 2)
    fields = dict(
        client=ClientInfo(nom="Client Test"),
        artisan=ArtisanInfo(nom="Artisan Test"),
        chantier=ChantierInfo(description="Chantier test"),
        lignes=lignes,
        totaux=TotauxDevis(total_ht=total_ht, total_tva=round(total_ht * 0.2, 2), total_ttc=round(total_ht * 1.2, 2)),
    )
    fields.update(overrides)
    return Devis(**fields)


def pdf_text(devis: Devis, document_type: str = "devis", with_tva: bool = True, document_date: str | None = None, **kwargs) -> str:
    """Génère le PDF et retourne tout le texte extrait (toutes pages concaténées)."""
    pdf_bytes = generate_quote_pdf(devis, document_type, with_tva, document_date, **kwargs)
    reader = PdfReader(io.BytesIO(pdf_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def docx_text(devis: Devis, document_type: str = "devis", with_tva: bool = True, document_date: str | None = None, **kwargs) -> str:
    """Génère le .docx et retourne tout le texte (paragraphes + cellules de tableaux)."""
    docx_bytes = generate_quote_docx(devis, document_type, with_tva, document_date, **kwargs)
    doc = Document(io.BytesIO(docx_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def make_signature_png_base64() -> str:
    """Une petite image PNG valide (trait simulant une signature), en base64 pur."""
    import base64
    from io import BytesIO
    from PIL import Image, ImageDraw

    img = Image.new("RGBA", (300, 100), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    draw.line([(10, 80), (80, 20), (150, 70), (220, 30), (290, 60)], fill=(20, 20, 20, 255), width=4)
    buf = BytesIO()
    img.save(buf, format="PNG")
    return base64.b64encode(buf.getvalue()).decode()
